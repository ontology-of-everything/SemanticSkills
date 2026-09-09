#!/usr/bin/env python3
"""html-slides · HTML deck → 截图贴图版 PPTX。

用法:
  python3 tools/export_pptx.py <deck>.html [-o out.pptx] [--notes <deck>.plan.md] [--scale 1|2]

每页 <section class="slide"> 用 Chromium 截 1920×1080 图铺满 16:9 幻灯片；不可编辑。
--notes 传入 plan.md 时，逐页表的「核心观点」列按 data-label 匹配写入备注。
依赖缺失时退出码 2 并打印安装命令；本脚本不自动安装。
"""
import argparse
import re
import sys
import tempfile
from pathlib import Path

W, H = 1920, 1080


def need(mod, pip, extra=""):
    try:
        return __import__(mod)
    except ImportError:
        sys.stderr.write(f"缺少依赖 {mod}。安装：pip install {pip}{extra}\n")
        sys.exit(2)


def parse_notes(plan_path):
    """从 plan.md 逐页表读 {label: 核心观点}。表头需含 label 与 核心观点 两列。"""
    notes = {}
    if not plan_path:
        return notes
    rows = [l for l in Path(plan_path).read_text(encoding="utf-8").splitlines() if l.startswith("|")]
    header = None
    for l in rows:
        cells = [c.strip() for c in l.strip().strip("|").split("|")]
        if header is None:
            if "label" in cells and "核心观点" in cells:
                header = cells
            continue
        if set("".join(cells)) <= set("-: "):
            continue
        if len(cells) == len(header):
            row = dict(zip(header, cells))
            notes[row["label"].strip("`")] = row["核心观点"]
    return notes


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("html")
    ap.add_argument("-o", "--out")
    ap.add_argument("--notes", help="<deck>.plan.md，核心观点写入备注")
    ap.add_argument("--scale", type=int, default=2, choices=(1, 2), help="截图倍率，默认 2")
    a = ap.parse_args()

    need("pptx", "python-pptx")
    need("playwright", "playwright", " && playwright install chromium")
    from playwright.sync_api import sync_playwright
    from pptx import Presentation
    from pptx.util import Inches

    src = Path(a.html).resolve()
    if not src.exists():
        sys.exit(f"找不到 {src}")
    out = Path(a.out) if a.out else src.with_suffix(".pptx")
    notes = parse_notes(a.notes)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    with sync_playwright() as p, tempfile.TemporaryDirectory() as tmp:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={"width": W, "height": H}, device_scale_factor=a.scale)
        page.goto(src.as_uri() + "?export")
        page.wait_for_load_state("networkidle")
        # 放映模式一页一屏；展开全部 build，layer 只截默认面板
        page.evaluate("document.body.classList.add('present'); document.querySelectorAll('.bar').forEach(b => b.remove())")
        labels = page.eval_on_selector_all(".slide", "els => els.map(e => e.dataset.label || '')")
        n = len(labels)
        if not n:
            sys.exit("未找到 <section class=\"slide\">")
        for i, label in enumerate(labels):
            page.evaluate(
                """i => {
                  document.querySelectorAll('.slide-wrap').forEach((w, j) => w.classList.toggle('cur', j === i));
                  document.querySelectorAll('.build').forEach(b => b.setAttribute('data-shown', ''));
                  document.documentElement.style.setProperty('--s', 1);
                }""",
                i,
            )
            page.wait_for_timeout(80)
            wrap = page.query_selector_all(".slide-wrap")[i]
            shot = Path(tmp) / f"{i + 1:03d}.png"
            wrap.screenshot(path=str(shot))
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(shot), 0, 0, prs.slide_width, prs.slide_height)
            txt = notes.get(label)
            if txt:
                slide.notes_slide.notes_text_frame.text = txt
            sys.stderr.write(f"\r{i + 1}/{n} {label[:30]:<30}")
        browser.close()

    prs.save(out)
    sys.stderr.write(f"\n写出 {out}（{n} 页，截图贴图版，不可编辑）\n")


if __name__ == "__main__":
    main()
